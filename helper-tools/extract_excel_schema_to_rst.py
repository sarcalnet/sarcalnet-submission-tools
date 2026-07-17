# The MIT License (MIT)
# Copyright (c) 2025 by the xcube team
#
# Permission is hereby granted, free of charge, to any person obtaining a
# copy of this software and associated documentation files (the "Software"),
# to deal in the Software without restriction, including without limitation
# the rights to use, copy, modify, merge, publish, distribute, sublicense,
# and/or sell copies of the Software, and to permit persons to whom the
# Software is furnished to do so, subject to the following conditions:
#
# The above copyright notice and this permission notice shall be included in
# all copies or substantial portions of the Software.
#
# THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
# IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
# FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
# AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
# LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING
# FROM, OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER
# DEALINGS IN THE SOFTWARE.

# python extract_excel_schema_to_rst.py # prints usage
# python extract_excel_schema_to_rst.py /path/to/template.xlsx


import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SKIP_SHEETS = ("README", "definitions")


def cell_text(value):
    if value is None:
        return ""
    return str(value)


def underline(title, char):
    return char * max(len(title), 3)


def literal_block(lines, text):
    if not text:
        return
    lines.append("::")
    lines.append("")
    for raw_line in text.splitlines():
        lines.append(f"   {raw_line}" if raw_line else "   ")


def field_block(lines, label, value):
    if value:
        lines.append(f":{label}: {value}")


def used_columns(ws, rows):
    columns = []
    for col in range(1, ws.max_column + 1):
        values = [ws.cell(row=row, column=col).value for row in rows]
        if any(value is not None for value in values):
            columns.append(col)
    return columns


def title_from_source(source):
    stem = source.stem
    stem = re.sub(r"_v\d+(?:\.\d+)*$", "", stem)
    stem = stem.replace("_submission_template", " Submission Template")
    words = stem.replace("_", " ").replace("-", " ").split()
    return " ".join("SARCalNet" if word.lower() == "sarcalnet" else word.title() for word in words)


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Generate an RST schema summary from an Excel workbook. By default, "
            "row 1 is used as the display title, row 2 as the field name, row 3 "
            "as the status, and rows 4-5 as literal comment blocks."
        )
    )
    parser.add_argument("source", type=Path, help="Input .xlsx workbook")
    parser.add_argument("--output", type=Path, help="Output .rst file. Defaults to <source-stem>.rst")
    parser.add_argument("--title", help="Document title. Defaults to a readable title from the source file name")
    parser.add_argument(
        "--skip-sheet",
        action="append",
        default=list(DEFAULT_SKIP_SHEETS),
        help="Sheet name to skip. Can be repeated. Defaults to README and definitions",
    )
    parser.add_argument("--display-row", type=int, default=1, help="Row containing human-readable column titles")
    parser.add_argument("--field-row", type=int, default=2, help="Row containing technical field names")
    parser.add_argument("--status-row", type=int, default=3, help="Row containing status values")
    parser.add_argument("--description-row", type=int, default=4, help="Row containing the description text")
    parser.add_argument("--details-row", type=int, default=5, help="Row containing additional information")
    return parser.parse_args()


def main():
    args = parse_args()
    source = args.source
    output = args.output or Path(f"{source.stem}.rst")
    doc_title = args.title or title_from_source(source)
    skip_sheets = {sheet.casefold() for sheet in args.skip_sheet}
    summary_rows = {
        args.display_row,
        args.field_row,
        args.status_row,
        args.description_row,
        args.details_row,
    }

    wb = load_workbook(source, read_only=True, data_only=False)

    lines = [
        doc_title,
        underline(doc_title, "="),
        "",
        f"Source file: ``{source.name}``",
        "",
    ]

    for ws in wb.worksheets:
        if ws.title.casefold() in skip_sheets:
            continue

        lines.extend([ws.title, underline(ws.title, "-"), ""])

        for col in used_columns(ws, summary_rows):
            display_name = cell_text(ws.cell(row=args.display_row, column=col).value)
            field_name = cell_text(ws.cell(row=args.field_row, column=col).value)
            status = cell_text(ws.cell(row=args.status_row, column=col).value)
            description = cell_text(ws.cell(row=args.description_row, column=col).value)
            details = cell_text(ws.cell(row=args.details_row, column=col).value)

            column_title = display_name or field_name or f"Column {col}"
            lines.extend([column_title, underline(column_title, "~")])
            field_block(lines, "Field name", field_name)
            field_block(lines, "Required", status)

            if description:
                lines.extend(["", "Description", "^^^^^^^^^^^", ""])
                literal_block(lines, description)

            if details:
                lines.extend(["", "Constraints", "^^^^^^^^^^^^^^^^^^^^^^", ""])
                literal_block(lines, details)

            lines.append("")

    output.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    print(output.resolve())


if __name__ == "__main__":
    main()
